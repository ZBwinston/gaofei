#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 代写/25、26 下的稿费 Excel 汇总成 App 可导入的 seed.json。

只读源文件，不改动任何 Excel。
输出：
  private/seed.json            App「导入备份」可直接读的备份文件（不提交 Git）
  private/import-report.txt    导入报告：条数、金额、异常明细

用法：
  python scripts/import_xlsx.py [--source 目录] [--out 目录]
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import hashlib
import json
import os
import sys
from collections import Counter, defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_SOURCE = os.path.join(os.path.dirname(ROOT), '代写')

PLATFORMS = [
    {'id': 'plat-fanmiao', 'name': '帆淼写作设计', 'color': '#b08a57', 'sort': 1},
    {'id': 'plat-xianzu', 'name': '仙租顺利通过代写服务', 'color': '#7a8b99', 'sort': 2},
    {'id': 'plat-lanling', 'name': '澜翎幻境商务设计', 'color': '#8c7e9b', 'sort': 3},
    {'id': 'plat-wechat', 'name': '微信', 'color': '#7e9b8a', 'sort': 4},
]
PLATFORM_BY_FILE = {
    '帆淼写作设计': 'plat-fanmiao',
    '仙租顺利通过代写服务': 'plat-xianzu',
    '澜翎幻境商务设计': 'plat-lanling',
    '微信': 'plat-wechat',
}
STATUS_MAP = {
    '完结': '已完成', '完成': '已完成', '已完结': '已完成', '已完成': '已完成', '已交付': '已完成',
    '已创建': '进行中', '进行中': '进行中', '写作中': '进行中', '制作中': '进行中',
    '待修改': '待修改', '修改': '待修改', '返修': '待修改',
    '待确认': '待确认',
    '取消': '已取消', '已取消': '已取消', '作废': '已取消',
    '未标记': '未标记',
}
NOTE_AS_STATUS = {'完结', '完成', '已完结', '已完成', '已创建', '进行中', '待修改', '取消', '已取消'}

EXPECTED_FILE = 'expected.json'
HEADER_ALIASES = {
    'ordered_at': ['订单时间', '下单时间', '日期'],
    'external_id': ['订单id', '订单ID', '客户id'],
    'order_no': ['订单号', '订单编号'],
    'customer_service': ['客服', '客服名称'],
    'title': ['作品', '作品名称', '标题'],
    'commission': ['接单佣金', '佣金'],
    'note': ['备注'],
    'received': ['到手金额', '实际到手'],
    'status': ['状态'],
    'bill': ['账单金额'],
}


def clean(value):
    if value is None:
        return ''
    if isinstance(value, str):
        return value.strip()
    return value


def to_date(value):
    if isinstance(value, dt.datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, dt.date):
        return value.strftime('%Y-%m-%d')
    text = str(value or '').strip()
    if not text:
        return ''
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d'):
        try:
            return dt.datetime.strptime(text[:10].replace('/', '-').replace('.', '-'), '%Y-%m-%d').strftime('%Y-%m-%d')
        except ValueError:
            pass
    import re
    matched = re.search(r'(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})', text)
    if matched:
        year, month, day = matched.groups()
        try:
            return dt.date(int(year), int(month), int(day)).strftime('%Y-%m-%d')
        except ValueError:
            return ''
    return ''


def to_amount(value):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    text = str(value).replace('¥', '').replace('￥', '').replace(',', '').strip()
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def record_id(platform_id, record):
    payload = '|'.join([
        platform_id,
        record['orderedAt'],
        record['orderNo'],
        record['externalId'],
        record['title'],
        f"{record['commission']:.2f}",
        '' if record['received'] is None else f"{record['received']:.2f}",
        record['note'],
        record['status'],
    ])
    return 'imp-' + hashlib.sha1(payload.encode('utf-8')).hexdigest()[:16]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', default=DEFAULT_SOURCE)
    parser.add_argument('--out', default=ROOT)
    args = parser.parse_args()

    files = sorted(glob.glob(os.path.join(args.source, '**', '*.xlsx'), recursive=True))
    if not files:
        sys.exit(f'没有找到 Excel：{args.source}')

    report = []
    anomalies = []
    merged_fills = 0
    note_as_status = 0
    records = []
    seen_ids = set()
    duplicates = 0
    per_file = []

    for path in files:
        relative = os.path.relpath(path, args.source).replace('\\', '/')
        parts = relative.split('/')
        year_folder = parts[0] if len(parts) >= 2 and parts[0].isdigit() else ''
        file_name = os.path.splitext(os.path.basename(path))[0]
        platform_id = PLATFORM_BY_FILE.get(file_name)
        if not platform_id:
            anomalies.append(f'跳过未知平台文件：{relative}')
            continue

        workbook = openpyxl.load_workbook(path, data_only=True)
        file_rows = 0
        for sheet in workbook.worksheets:
            header_cells = [clean(cell.value) for cell in sheet[1]]
            columns = {}
            for key, aliases in HEADER_ALIASES.items():
                for index, header in enumerate(header_cells):
                    if header in aliases:
                        columns[key] = index
                        break
            if 'ordered_at' not in columns:
                anomalies.append(f'跳过无「订单时间」的表：{relative} / {sheet.title}')
                continue

            month_match = None
            import re as _re
            month_match = _re.search(r'(\d{1,2})月', str(sheet.title))
            sheet_month = int(month_match.group(1)) if month_match and 1 <= int(month_match.group(1)) <= 12 else None
            month_mismatch = 0

            merged_values = {}
            for merged in sheet.merged_cells.ranges:
                top_left = sheet.cell(merged.min_row, merged.min_col).value
                for row in range(merged.min_row, merged.max_row + 1):
                    for col in range(merged.min_col, merged.max_col + 1):
                        if (row, col) != (merged.min_row, merged.min_col):
                            merged_values[(row, col)] = top_left

            for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                cells = []
                for col_index, cell in enumerate(row, start=1):
                    value = cell.value
                    if (value is None or value == '') and (row_index, col_index) in merged_values:
                        value = merged_values[(row_index, col_index)]
                        merged_fills += 1
                    cells.append(value)
                if not any(clean(value) != '' for value in cells):
                    continue

                def get(key):
                    index = columns.get(key)
                    if index is None or index >= len(cells):
                        return ''
                    return cells[index]

                ordered_at = to_date(get('ordered_at'))
                if not ordered_at:
                    anomalies.append(f'{relative} / {sheet.title} 第{row_index}行：日期无法识别，已跳过')
                    continue

                note = str(clean(get('note')) or '')
                status_raw = str(clean(get('status')) or '')
                if not status_raw and note in NOTE_AS_STATUS:
                    status_raw, note = note, ''
                    note_as_status += 1
                status = STATUS_MAP.get(status_raw, '未标记' if not status_raw else '待确认')

                commission = to_amount(get('commission'))
                if commission is None:
                    commission = to_amount(get('bill'))
                if commission is None:
                    commission = 0

                received = to_amount(get('received'))
                inferred = False
                if received is None and status == '已完成' and not note.strip():
                    received = commission
                    inferred = True

                record = {
                    'id': '',
                    'platformId': platform_id,
                    'orderedAt': ordered_at,
                    'orderNo': str(clean(get('order_no')) or ''),
                    'externalId': str(clean(get('external_id')) or ''),
                    'customerService': str(clean(get('customer_service')) or ''),
                    'title': str(clean(get('title')) or ''),
                    'commission': round(float(commission), 2),
                    'received': None if received is None else round(float(received), 2),
                    'receivedInferred': inferred,
                    'note': note,
                    'status': status,
                    'createdAt': f'{ordered_at}T00:00:00.000Z',
                    'updatedAt': f'{ordered_at}T00:00:00.000Z',
                    'deletedAt': None,
                    'source': {
                        'year': year_folder,
                        'file': file_name,
                        'sheet': sheet.title,
                        'row': row_index,
                    },
                }
                record['id'] = record_id(platform_id, record)
                if record['id'] in seen_ids:
                    duplicates += 1
                    record['id'] = f"{record['id']}-{duplicates}"
                seen_ids.add(record['id'])
                records.append(record)
                file_rows += 1

                if sheet_month and int(record['orderedAt'][5:7]) != sheet_month:
                    month_mismatch += 1
        if month_mismatch:
            anomalies.append(
                f'{relative} / {sheet.title}：{month_mismatch} 行日期月份与表名不同（已按日期归期）'
            )
        per_file.append((relative, file_rows))

    records.sort(key=lambda item: (item['orderedAt'], item['platformId'], item['title']))

    commission_total = round(sum(item['commission'] for item in records), 2)
    received_total = round(sum(item['received'] or 0 for item in records), 2)
    diff_total = round(sum(
        (item['commission'] or 0) - (item['received'] or 0)
        for item in records if item['received'] is not None
    ), 2)
    inferred_count = sum(1 for item in records if item['receivedInferred'])
    pending = [item for item in records if item['received'] is None]
    status_counts = Counter(item['status'] for item in records)

    by_platform = defaultdict(lambda: {'count': 0, 'commission': 0.0, 'received': 0.0})
    for item in records:
        bucket = by_platform[item['platformId']]
        bucket['count'] += 1
        bucket['commission'] += item['commission']
        bucket['received'] += item['received'] or 0

    actual = {
        'records': len(records),
        'platforms': len({item['platformId'] for item in records}),
        'commission': commission_total,
        'received': received_total,
        'diff': diff_total,
        'inferred': inferred_count,
        'pending': len(pending),
    }
    expected_path = os.path.join(args.out, 'private', EXPECTED_FILE)
    expected = None
    if os.path.exists(expected_path):
        with open(expected_path, encoding='utf-8') as handle:
            expected = json.load(handle)

    mismatches = {}
    if expected:
        for key, expected_value in expected.items():
            if key in actual and abs(expected_value - actual[key]) > 0.004:
                mismatches[key] = (expected_value, actual[key])
    if mismatches:
        print('!! 口径与预期不一致：', mismatches)
        sys.exit(2)

    seed = {
        'app': '稿费记录',
        'version': 1,
        'exportedAt': dt.datetime.now().astimezone().isoformat(timespec='seconds'),
        'platforms': PLATFORMS,
        'records': records,
        'settings': {'theme': 'system', 'onboarded': False, 'lastBackupAt': ''},
        'sync': {'lastPulledServerAt': '', 'lastSyncedAt': '', 'dirty': [], 'userId': '', 'email': ''},
    }

    out_dir = args.out
    os.makedirs(os.path.join(out_dir, 'private'), exist_ok=True)
    seed_path = os.path.join(out_dir, 'private', 'seed.json')
    with open(seed_path, 'w', encoding='utf-8') as handle:
        json.dump(seed, handle, ensure_ascii=False, indent=2)
        handle.write('\n')

    lines = []
    lines.append('稿费记录 · 导入报告')
    lines.append(f'生成时间：{dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    lines.append(f'源目录：{args.source}')
    lines.append('')
    lines.append('== 文件 ==')
    for name, count in per_file:
        lines.append(f'  {name}: {count} 条')
    lines.append('')
    lines.append('== 口径 ==')
    lines.append(f'  记录数：{len(records)}')
    lines.append(f'  平台数：{actual["platforms"]}')
    lines.append(f'  接单佣金合计：{commission_total:.2f}')
    lines.append(f'  到手合计：{received_total:.2f}（其中自动推定 {inferred_count} 条）')
    lines.append(f'  差额合计：{diff_total:.2f}')
    lines.append(f'  仍未填到手：{len(pending)} 条，佣金 {sum(item["commission"] for item in pending):.2f}')
    lines.append('  状态分布：' + '，'.join(f'{key} {value}' for key, value in sorted(status_counts.items())))
    lines.append('')
    lines.append('== 分平台 ==')
    for platform in PLATFORMS:
        bucket = by_platform[platform['id']]
        lines.append(
            f'  {platform["name"]}：{bucket["count"]} 条，佣金 {bucket["commission"]:.2f}，到手 {bucket["received"]:.2f}'
        )
    lines.append('')
    lines.append('== 处理 ==')
    lines.append(f'  合并单元格补全：{merged_fills} 个单元格')
    lines.append(f'  备注列误填状态并归位：{note_as_status} 行')
    lines.append(f'  内容重复并按序编号：{duplicates} 行')
    lines.append('')
    lines.append('== 仍未填到手（请人工确认） ==')
    for item in pending:
        lines.append(
            f'  {item["orderedAt"]} {platform_name(item["platformId"])} 佣金 {item["commission"]:.2f} '
            f'状态 {item["status"]} 备注「{item["note"]}」'
        )
    lines.append('')
    lines.append('== 异常与提示 ==')
    if anomalies:
        for item in anomalies:
            lines.append(f'  - {item}')
    else:
        lines.append('  无')
    lines.append('')
    lines.append('== 校验 ==')
    if expected:
        lines.append('  与 private/expected.json 的固定口径一致。' if not mismatches else f'  与预期不一致：{mismatches}')
    else:
        lines.append('  未找到 private/expected.json，只生成了统计结果。')

    report_dir = os.path.join(out_dir, 'private')
    report_path = os.path.join(report_dir, 'import-report.txt')
    with open(report_path, 'w', encoding='utf-8') as handle:
        handle.write('\n'.join(lines) + '\n')

    print(f'OK 记录 {len(records)} 条，佣金 {commission_total:.2f}，到手 {received_total:.2f}，差额 {diff_total:.2f}')
    print(f'seed: {seed_path}')
    print(f'报告: {report_path}')


def platform_name(platform_id):
    for item in PLATFORMS:
        if item['id'] == platform_id:
            return item['name']
    return platform_id


if __name__ == '__main__':
    main()
